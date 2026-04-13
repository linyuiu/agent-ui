import json
from typing import Any, Dict, List, Optional

import openpyxl
import requests


# =========================
# 配置区
# =========================
EXCEL_FILE = "/Users/linyu/Desktop/IVD产品竞品信息.xlsx"

API_URL = "https://ai.mindrayanimal.com/admin/api/workspace/default/knowledge/019d47f6-383b-7482-87c5-8a1007258bd5/document/batch_create"


AUTHORIZATION = "Bearer user-f86e223427899afdc5b711ba498e8bdc"

# 如果接口需要 source_file_id，就填；否则设为 None
SOURCE_FILE_ID = "019d4d63-9aa9-76a2-8f94-743695af7240"

# 仅处理指定 sheet；None 表示处理全部
TARGET_SHEETS = None  # 例如 ["生化免疫", "血球"]

HEADER_ROW = 1
MODEL_COLUMN_NAME = "产品型号"
EMPTY_VALUE = ""
REQUEST_TIMEOUT = 60


# =========================
# Excel 处理
# =========================
def normalize_cell_value(value: Any) -> str:
    if value is None:
        return EMPTY_VALUE

    text = str(value).strip()
    if not text:
        return EMPTY_VALUE

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\n", "<br>")
    text = text.replace("|", "\\|")
    return text


def expand_merged_cells(ws) -> None:
    """
    把工作表中的合并单元格展开，
    将左上角值填充到整个合并区域。
    """
    merged_ranges = list(ws.merged_cells.ranges)

    for merged_range in merged_ranges:
        min_row, max_row = merged_range.min_row, merged_range.max_row
        min_col, max_col = merged_range.min_col, merged_range.max_col

        top_left_value = ws.cell(min_row, min_col).value

        ws.unmerge_cells(str(merged_range))

        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(r, c).value = top_left_value


def get_headers(ws, header_row: int) -> List[str]:
    headers = []
    for c in range(1, ws.max_column + 1):
        value = ws.cell(header_row, c).value
        headers.append(normalize_cell_value(value) or f"列{c}")
    return headers


def sheet_to_rows(ws, header_row: int = 1) -> List[Dict[str, str]]:
    headers = get_headers(ws, header_row)
    rows: List[Dict[str, str]] = []

    for r in range(header_row + 1, ws.max_row + 1):
        row_dict: Dict[str, str] = {}
        has_any_value = False

        for c in range(1, ws.max_column + 1):
            value = normalize_cell_value(ws.cell(r, c).value)
            row_dict[headers[c - 1]] = value
            if value != EMPTY_VALUE:
                has_any_value = True

        if has_any_value:
            rows.append(row_dict)

    return rows


def group_rows_by_model(rows: List[Dict[str, str]], model_column_name: str) -> Dict[str, List[Dict[str, str]]]:
    grouped: Dict[str, List[Dict[str, str]]] = {}

    for row in rows:
        model = row.get(model_column_name, "").strip()
        if not model:
            model = "未命名产品型号"
        grouped.setdefault(model, []).append(row)

    return grouped


def build_markdown_table(rows: List[Dict[str, str]]) -> str:
    if not rows:
        return ""

    headers = list(rows[0].keys())
    header_line = "| " + " | ".join(headers) + " |"
    separator_line = "| " + " | ".join(["---"] * len(headers)) + " |"

    body_lines = []
    for row in rows:
        values = [row.get(h, EMPTY_VALUE) for h in headers]
        body_lines.append("| " + " | ".join(values) + " |")

    return "\n".join([header_line, separator_line] + body_lines)


def build_payload(
    excel_file: str,
    header_row: int = 1,
    model_column_name: str = "产品型号",
    target_sheets: Optional[List[str]] = None,
    source_file_id: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """
    构造最终请求体：
    [
      {
        "name": "血球",
        "paragraphs": [
          {"title": "", "content": "...某型号markdown表格..."},
          {"title": "", "content": "...另一个型号markdown表格..."}
        ],
        "source_file_id": "xxx"
      },
      {
        "name": "生化免疫",
        "paragraphs": [...]
      }
    ]
    """
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    payload: List[Dict[str, Any]] = []

    worksheets = wb.worksheets
    if target_sheets:
        worksheets = [wb[s] for s in target_sheets if s in wb.sheetnames]

    for ws in worksheets:
        expand_merged_cells(ws)

        rows = sheet_to_rows(ws, header_row=header_row)
        if not rows:
            continue

        headers = list(rows[0].keys())
        if model_column_name not in headers:
            raise ValueError(f"Sheet [{ws.title}] 未找到列：{model_column_name}，当前表头：{headers}")

        grouped = group_rows_by_model(rows, model_column_name)

        paragraphs = []
        for model, model_rows in grouped.items():
            markdown = build_markdown_table(model_rows)

            paragraphs.append({
                "title": "",
                "content": markdown
            })

        doc = {
            "name": ws.title,
            "paragraphs": paragraphs
        }

        if source_file_id:
            doc["source_file_id"] = source_file_id

        payload.append(doc)

    return payload


def send_payload(api_url: str, authorization: str, payload: List[Dict[str, Any]]) -> requests.Response:
    headers = {
        "Authorization": authorization,
        "Content-Type": "application/json"
    }

    response = requests.put(
        api_url,
        headers=headers,
        json=payload,
        timeout=REQUEST_TIMEOUT
    )
    return response


def main():
    payload = build_payload(
        excel_file=EXCEL_FILE,
        header_row=HEADER_ROW,
        model_column_name=MODEL_COLUMN_NAME,
        target_sheets=TARGET_SHEETS,
        source_file_id=SOURCE_FILE_ID,
    )



    # 只调用一次接口
    response = send_payload(
        api_url=API_URL,
        authorization=AUTHORIZATION,
        payload=payload
    )




if __name__ == "__main__":
    main()